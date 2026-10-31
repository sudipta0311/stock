"""Official current constituent files for US and Indian broad-market universes."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from typing import Any

from wealth.marketdata.sources.http import get
from wealth.marketdata.sources.nse import NseSecurity

SP500_HOLDINGS_URL = (
    "https://www.ssga.com/us/en/institutional/library-content/products/fund-data/etfs/us/"
    "holdings-daily-us-en-spy.xlsx"
)
SPY_NAV_HISTORY_URL = (
    "https://www.ssga.com/us/en/intermediary/library-content/products/fund-data/etfs/us/"
    "navhist-us-en-spy.xlsx"
)
SP500_SECTORS_URL = "https://en.wikipedia.org/wiki/List_of_S%26P_500_companies"
NIFTY500_ARCHIVE_URL = "https://archives.nseindia.com/content/indices/ind_nifty500list.csv"
SP500_SOURCE = "SSGA_SPY_HOLDINGS"
SP500_PRICE_SOURCE = "SSGA_SPY_IMPLIED"
NIFTY500_SOURCE = "NSE_NIFTY500_ARCHIVE"
SURVIVORSHIP_LABEL = "current constituents only (survivorship bias)"
# US tickers start with a letter; SPY's workbook also lists "-" and numeric placeholder codes.
_TICKER_PATTERN = re.compile(r"^[A-Z][A-Z0-9.-]{0,14}$")


@dataclass(frozen=True)
class UniverseSecurity:
    symbol: str
    name: str
    isin: str | None = None
    series: str | None = None
    sector: str = "Unclassified"
    industry: str = "Unclassified"
    weight_pct: Decimal | None = None
    shares_held: Decimal | None = None


@dataclass(frozen=True)
class Sp500Holdings:
    securities: list[UniverseSecurity]
    as_of: date


@dataclass(frozen=True)
class Sp500Sector:
    sector: str
    industry: str


def _key(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _field(row: dict[str, Any], *names: str) -> str | None:
    normalized = {_key(str(key)): value for key, value in row.items() if key is not None}
    for name in names:
        value = normalized.get(_key(name))
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _classification_label(value: str | None, limit: int) -> str:
    if value is None or value.casefold() in {"-", "n/a", "na", "none", "null", "unknown"}:
        return "Unclassified"
    return value[:limit] or "Unclassified"


def _decimal(value: Any) -> Decimal | None:
    try:
        parsed = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None
    return parsed if parsed.is_finite() else None


def _parse_sp500_as_of(value: Any) -> date | None:
    if not isinstance(value, str):
        return None
    cleaned = value.strip().removeprefix("As of ").strip()
    for pattern in ("%d-%b-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(cleaned, pattern).date()
        except ValueError:
            continue
    return None


def parse_sp500_holdings(content: bytes) -> Sp500Holdings:
    """Read SPY's daily issuer workbook and exclude cash/unassigned positions."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        metadata: dict[str, Any] = {}
        header: list[str] | None = None
        holdings: list[UniverseSecurity] = []
        for row in rows:
            if not any(value is not None for value in row):
                continue
            first = str(row[0] or "").strip().casefold()
            if first in {"fund name:", "ticker symbol:", "holdings:"}:
                metadata[first.rstrip(":")] = row[1] if len(row) > 1 else None
                continue
            values = [str(value).strip() if value is not None else "" for value in row]
            if header is None:
                if "ticker" in {_key(value) for value in values}:
                    header = values
                continue
            raw = dict(zip(header, row, strict=False))
            symbol = _field(raw, "Ticker", "Symbol")
            name = _field(raw, "Name", "Security Name")
            sector = (_field(raw, "Sector") or "").casefold()
            sector_label = _classification_label(_field(raw, "Sector"), 80)
            industry_label = _classification_label(_field(raw, "Industry", "Industry Name"), 120)
            if not symbol or not name or sector == "unassigned":
                continue
            symbol = symbol.upper()
            if not _TICKER_PATTERN.fullmatch(symbol):
                continue
            holdings.append(
                UniverseSecurity(
                    symbol=symbol,
                    name=name[:160],
                    sector=sector_label,
                    industry=industry_label,
                    weight_pct=_decimal(_field(raw, "Weight")),
                    shares_held=_decimal(_field(raw, "Shares Held")),
                )
            )
        as_of = _parse_sp500_as_of(metadata.get("holdings"))
        if as_of is None:
            raise ValueError("SPY holdings workbook has no recognized as-of date")
        unique = {item.symbol: item for item in holdings}
        if not unique:
            raise ValueError("SPY holdings workbook has no equity constituents")
        return Sp500Holdings(
            securities=sorted(unique.values(), key=lambda item: item.symbol),
            as_of=as_of,
        )
    finally:
        workbook.close()


def fetch_sp500() -> Sp500Holdings:
    return parse_sp500_holdings(get(SP500_HOLDINGS_URL).content)


def parse_spy_total_net_assets(content: bytes) -> dict[date, Decimal]:
    """Read SPY's daily NAV history: total net assets by date."""
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        header: list[str] | None = None
        result: dict[date, Decimal] = {}
        for row in workbook.active.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if header is None:
                if {"date", "totalnetassets"} <= {_key(value) for value in values}:
                    header = values
                continue
            raw = dict(zip(header, row, strict=False))
            on = _parse_sp500_as_of(_field(raw, "Date"))
            assets = _decimal(_field(raw, "Total Net Assets"))
            if on is not None and assets is not None and assets > 0:
                result[on] = assets
        if not result:
            raise ValueError("SPY NAV history has no total net assets")
        return result
    finally:
        workbook.close()


def implied_closes(
    holdings: Sp500Holdings, total_net_assets: dict[date, Decimal]
) -> dict[str, Decimal]:
    """Each holding's close on the holdings date: weight x fund net assets / shares held.

    SSGA's weights are each position's share of the fund's net assets at that day's close,
    so the issuer's two files give a price without a separate market-data feed.
    """
    assets = total_net_assets.get(holdings.as_of)
    if assets is None:
        return {}
    closes: dict[str, Decimal] = {}
    for item in holdings.securities:
        if not item.weight_pct or not item.shares_held or item.shares_held <= 0:
            continue
        close = item.weight_pct / Decimal(100) * assets / item.shares_held
        if close > 0:
            closes[item.symbol] = close.quantize(Decimal("0.000001"))
    return closes


def fetch_sp500_closes(holdings: Sp500Holdings) -> dict[str, Decimal]:
    return implied_closes(holdings, parse_spy_total_net_assets(get(SPY_NAV_HISTORY_URL).content))


class _ConstituentsTable(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.depth = 0
        self.rows: list[list[str]] = []
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "table":
            if self.depth or dict(attrs).get("id") == "constituents":
                self.depth += 1
        elif self.depth == 1 and tag == "tr":
            self.rows.append([])
        elif self.depth == 1 and tag in {"td", "th"} and self.rows:
            self._cell = []

    def handle_endtag(self, tag: str) -> None:
        if tag == "table" and self.depth:
            self.depth -= 1
        elif self.depth == 1 and tag in {"td", "th"} and self._cell is not None:
            self.rows[-1].append(" ".join("".join(self._cell).split()))
            self._cell = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def parse_sp500_sectors(html: str) -> dict[str, Sp500Sector]:
    """Read GICS sectors by ticker from the public S&P 500 constituents table."""
    table = _ConstituentsTable()
    table.feed(html)
    if not table.rows:
        raise ValueError("S&P 500 constituents table not found")
    header = [_key(cell) for cell in table.rows[0]]
    try:
        symbol_at = header.index("symbol")
        sector_at = header.index("gicssector")
        industry_at = header.index("gicssubindustry")
    except ValueError as error:
        raise ValueError("S&P 500 constituents table has unexpected columns") from error
    result: dict[str, Sp500Sector] = {}
    for row in table.rows[1:]:
        if len(row) <= max(symbol_at, sector_at, industry_at):
            continue
        symbol = row[symbol_at].upper().replace("-", ".")
        if not _TICKER_PATTERN.fullmatch(symbol):
            continue
        result[symbol] = Sp500Sector(
            sector=_classification_label(row[sector_at] or None, 80),
            industry=_classification_label(row[industry_at] or None, 120),
        )
    if not result:
        raise ValueError("S&P 500 constituents table has no rows")
    return result


def fetch_sp500_sectors() -> dict[str, Sp500Sector]:
    return parse_sp500_sectors(get(SP500_SECTORS_URL).text)


def parse_nifty500_csv(text: str) -> list[NseSecurity]:
    """Parse NSE's public index archive and deduplicate its current tickers."""
    reader = csv.DictReader(io.StringIO(text.lstrip("\ufeff")))
    result: dict[str, NseSecurity] = {}
    for row in reader:
        symbol = _field(row, "Symbol", "Ticker")
        name = _field(row, "Company Name", "Name")
        isin = _field(row, "ISIN Code", "ISIN")
        series = _field(row, "Series")
        if not symbol or not name:
            continue
        symbol = symbol.upper()
        isin_value = isin.upper() if isin and isin != "-" else None
        series_value = series.upper() if series and series != "-" else None
        industry = _classification_label(
            _field(row, "Industry", "Industry Name", "Basic Industry"), 120
        )
        sector = _classification_label(_field(row, "Sector"), 80)
        if sector == "Unclassified" and industry != "Unclassified":
            sector = industry[:80]
        result[symbol] = NseSecurity(
            symbol=symbol,
            name=name[:160],
            isin=isin_value,
            series=series_value,
            security_type="equity",
            sleeve="equity_india",
            category="equity:india",
            sector=sector,
            industry=industry,
        )
    if not result:
        raise ValueError("NSE NIFTY 500 archive CSV has no constituents")
    return sorted(result.values(), key=lambda item: item.symbol)


def fetch_nifty500() -> list[NseSecurity]:
    return parse_nifty500_csv(get(NIFTY500_ARCHIVE_URL).text)
