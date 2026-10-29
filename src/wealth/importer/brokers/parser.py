"""Header-driven CSV and XLSX broker statement parsing."""

from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
from collections.abc import Mapping
from contextlib import suppress
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from wealth.core.dates import add_months
from wealth.core.money import parse_inr
from wealth.importer.brokers.types import BrokerHolding, Layout, Trade

_ALIASES: dict[str, tuple[str, ...]] = {
    "isin": ("isin", "isin code"),
    "symbol": ("symbol", "trading symbol", "ticker", "scrip", "stock", "instrument"),
    "name": ("name", "company name", "stock name", "security name", "instrument"),
    "quantity": ("quantity", "qty", "shares", "units", "holdings"),
    "avg_price": (
        "avg price",
        "average price",
        "average cost value",
        "buy price",
        "cost price",
        "avg cost",
        "purchase price",
        "acquisition price",
        "avg buy price",
    ),
    "last_price": ("ltp", "last price", "current price", "market price", "close price", "cmp"),
    "exchange": ("exchange",),
    "trade_date": ("trade date", "date", "transaction date"),
    "trade_type": ("trade type", "type", "action", "transaction type"),
    "price": ("price", "trade price", "rate"),
    "order_id": ("order id", "order number"),
    "trade_id": ("trade id", "trade no", "trade number", "execution id"),
}
_KNOWN_HEADERS = frozenset(alias for values in _ALIASES.values() for alias in values)
_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d-%b-%Y")


def normalize_header(value: object) -> str:
    """Normalize a vendor header and discard bracketed units such as currency labels."""
    text = unicodedata.normalize("NFKC", str(value or "")).casefold().strip()
    text = re.sub(r"\([^)]*\)|\[[^]]*\]", " ", text)
    text = text.replace("₨", " ").replace("₹", " ")
    text = re.sub(r"[._]+", " ", text)
    text = re.sub(r"[^\w\s-]", " ", text)
    text = re.sub(r"[_\s]+", " ", text)
    return text.strip()


def header_fingerprint(headers: list[str] | tuple[str, ...]) -> str:
    normalized = sorted(
        {normalize_header(header) for header in headers if normalize_header(header)}
    )
    return hashlib.sha256("\n".join(normalized).encode("utf-8")).hexdigest()


def load_layouts(path: Path | None = None) -> tuple[Layout, ...]:
    source = path or Path(__file__).resolve().parents[4] / "rules" / "brokers.yaml"
    document = yaml.safe_load(source.read_text(encoding="utf-8"))
    result: list[Layout] = []
    for entry in document.get("layouts", []):
        columns = {
            str(field): tuple(str(alias) for alias in aliases)
            for field, aliases in entry.get("columns", {}).items()
        }
        result.append(
            Layout(
                broker=str(entry["broker"]),
                kind=entry["kind"],
                fingerprint=frozenset(normalize_header(v) for v in entry["fingerprint"]),
                columns=columns,
                verified=bool(entry.get("verified", True)),
                source_note=str(entry.get("source_note", "")),
            )
        )
    return tuple(result)


def _layout_from_template(template: Mapping[str, Any]) -> Layout:
    columns = {
        str(field): str(header)
        for field, header in dict(template.get("columns") or {}).items()
        if header is not None
    }
    return Layout(
        broker=str(template.get("name") or "saved template"),
        kind=template["kind"],
        fingerprint=frozenset(str(value) for value in template["header_names"]),
        columns=columns,
        source_note="Saved by this household.",
        header_fingerprint=str(template["header_fingerprint"]),
    )


def detect_layout(
    headers: list[str] | tuple[str, ...],
    layouts: tuple[Layout, ...] | list[Layout] | None = None,
    templates: tuple[Mapping[str, Any], ...] | list[Mapping[str, Any]] = (),
) -> Layout | None:
    normalized = {normalize_header(header) for header in headers}
    fingerprint = header_fingerprint(headers)
    for template in templates:
        if template.get("header_fingerprint") == fingerprint:
            return _layout_from_template(template)
    for layout in layouts if layouts is not None else load_layouts():
        if layout.fingerprint <= normalized:
            return layout
    return None


def suggested_mapping(headers: list[str] | tuple[str, ...]) -> dict[str, str | None]:
    by_normalized = {normalize_header(header): header for header in headers}
    result: dict[str, str | None] = {}
    for canonical, aliases in _ALIASES.items():
        result[canonical] = next(
            (by_normalized[alias] for alias in aliases if alias in by_normalized), None
        )
    return result


def _header_row(rows: list[tuple[Any, ...]]) -> int | None:
    for index, row in enumerate(rows):
        non_empty = [value for value in row if value is not None and str(value).strip()]
        if len(non_empty) >= 3 and any(
            normalize_header(value) in _KNOWN_HEADERS for value in non_empty
        ):
            return index
    return None


def read_table(data: bytes, filename: str) -> tuple[list[str], list[dict[str, Any]]]:
    """Read a CSV or first suitable XLSX worksheet without evaluating formulas."""
    suffix = Path(filename).suffix.casefold()
    if suffix == ".xlsx" and data.startswith(b"PK"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for worksheet in workbook.worksheets:
                values = list(worksheet.iter_rows(values_only=True))
                row_index = _header_row(values)
                if row_index is None:
                    continue
                return _rows_from_values(values, row_index)
        finally:
            workbook.close()
        raise ValueError("No worksheet contains a recognizable header row.")
    if suffix != ".csv" or data.startswith(b"PK"):
        raise ValueError("The file does not match a supported CSV or XLSX format.")

    decoded: str | None = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            decoded = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if decoded is None:
        raise ValueError("The CSV encoding is not supported.")
    rows = list(csv.reader(io.StringIO(decoded, newline="")))
    row_index = _header_row([tuple(row) for row in rows])
    if row_index is None:
        if rows and len([cell for cell in rows[0] if cell.strip()]) >= 3:
            row_index = 0
        else:
            raise ValueError("No recognizable header row was found.")
    return _rows_from_values([tuple(row) for row in rows], row_index)


def _rows_from_values(
    values: list[tuple[Any, ...]], row_index: int
) -> tuple[list[str], list[dict[str, Any]]]:
    raw_headers = list(values[row_index])
    while raw_headers and (raw_headers[-1] is None or not str(raw_headers[-1]).strip()):
        raw_headers.pop()
    headers = [str(value or "").strip() for value in raw_headers]
    rows: list[dict[str, Any]] = []
    for row in values[row_index + 1 :]:
        if not row or not any(value is not None and str(value).strip() for value in row):
            continue
        rows.append(
            {
                headers[index]: row[index] if index < len(row) else None
                for index in range(len(headers))
            }
        )
    return headers, rows


def _column_value(row: Mapping[str, Any], layout: Layout, canonical: str) -> Any:
    configured = layout.columns.get(canonical)
    if configured is None:
        return None
    aliases = (configured,) if isinstance(configured, str) else configured
    normalized = {normalize_header(header): value for header, value in row.items()}
    for alias in aliases:
        key = normalize_header(alias)
        if key in normalized:
            return normalized[key]
    return None


def _decimal_value(value: Any, *, field: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{field} must be a number")
    raw = str(value).strip().replace("₹", "").replace("₨", "").replace("INR", "")
    raw = raw.replace("Rs.", "").replace("rs.", "").strip()
    if not raw:
        raise ValueError(f"{field} must be a number")
    try:
        # The shared C02 parser validates Indian comma grouping and rupee precision.
        return parse_inr(raw)
    except ValueError:
        cleaned = raw.replace(",", "")
        if not re.fullmatch(r"-?\d+(?:\.\d+)?", cleaned):
            raise ValueError(f"{field} must be a number") from None
        try:
            number = Decimal(cleaned)
        except InvalidOperation as error:
            raise ValueError(f"{field} must be a number") from error
        if not number.is_finite():
            raise ValueError(f"{field} must be finite") from None
        return number


def _date_value(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in _DATE_FORMATS:
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise ValueError("trade date must use a supported date format")


def _text(value: Any) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


def _is_summary(row: Mapping[str, Any], layout: Layout) -> bool:
    values = [str(value).strip().casefold() for value in row.values() if value is not None]
    return not any(values) or any(
        value in {"total", "grand total", "summary", "subtotal"} for value in values
    )


def parse_holdings(rows: list[dict[str, Any]], layout: Layout) -> list[BrokerHolding]:
    if layout.kind != "holdings":
        raise ValueError("A holdings layout is required.")
    result: list[BrokerHolding] = []
    for row in rows:
        if _is_summary(row, layout):
            continue
        try:
            quantity = _decimal_value(_column_value(row, layout, "quantity"), field="quantity")
            avg_price = _decimal_value(
                _column_value(row, layout, "avg_price"), field="average price"
            )
        except ValueError:
            continue
        if quantity <= 0 or avg_price < 0:
            continue
        symbol = _text(_column_value(row, layout, "symbol"))
        name = _text(_column_value(row, layout, "name"))
        if (symbol or name or _text(_column_value(row, layout, "isin"))) is None:
            continue
        last_price = None
        raw_last = _column_value(row, layout, "last_price")
        if raw_last not in (None, ""):
            with suppress(ValueError):
                last_price = _decimal_value(raw_last, field="last price")
        isin = _text(_column_value(row, layout, "isin"))
        result.append(
            BrokerHolding(
                isin=isin.upper() if isin else None,
                symbol=symbol.upper() if symbol else None,
                name=name,
                quantity=quantity,
                avg_price=avg_price,
                last_price=last_price,
                exchange=_text(_column_value(row, layout, "exchange")),
            )
        )
    return result


def parse_tradebook(rows: list[dict[str, Any]], layout: Layout) -> list[Trade]:
    if layout.kind != "tradebook":
        raise ValueError("A tradebook layout is required.")
    result: list[Trade] = []
    for row in rows:
        if _is_summary(row, layout):
            continue
        raw_type = str(_column_value(row, layout, "trade_type") or "").strip().casefold()
        if raw_type in {"buy", "b", "purchase"}:
            trade_type = "buy"
        elif raw_type in {"sell", "s", "sale"}:
            trade_type = "sell"
        else:
            continue
        try:
            trade_date = _date_value(_column_value(row, layout, "trade_date"))
            quantity = _decimal_value(_column_value(row, layout, "quantity"), field="quantity")
            price = _decimal_value(_column_value(row, layout, "price"), field="price")
        except ValueError:
            continue
        if quantity <= 0 or price < 0:
            continue
        isin = _text(_column_value(row, layout, "isin"))
        symbol = _text(_column_value(row, layout, "symbol"))
        if isin is None and symbol is None:
            continue
        result.append(
            Trade(
                isin=isin.upper() if isin else None,
                symbol=symbol.upper() if symbol else None,
                trade_date=trade_date,
                trade_type=trade_type,
                quantity=quantity,
                price=price,
                exchange=_text(_column_value(row, layout, "exchange")),
                order_id=_text(_column_value(row, layout, "order_id")),
                trade_id=_text(_column_value(row, layout, "trade_id")),
            )
        )
    return result


def is_long_term(acquired_on: date, sold_on: date) -> bool:
    """Listed equity is long-term after twelve calendar months."""
    return sold_on > add_months(acquired_on, 12)
